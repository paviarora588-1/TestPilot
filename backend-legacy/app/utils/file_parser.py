import csv
import mimetypes
from pathlib import Path
from typing import Any


SUPPORTED_IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp"}


def parse_uploaded_file(file_path: str) -> dict[str, Any]:
    path = Path(file_path)
    suffix = path.suffix.lower()

    if suffix == ".txt":
        text = path.read_text(encoding="utf-8", errors="ignore")
    elif suffix == ".csv":
        text = _parse_csv(path)
    elif suffix == ".xlsx":
        text = _parse_xlsx(path)
    elif suffix == ".pdf":
        text = _parse_pdf(path)
    elif suffix == ".docx":
        text = _parse_docx(path)
    elif suffix in SUPPORTED_IMAGE_EXTENSIONS:
        mime_type, _ = mimetypes.guess_type(str(path))
        text = f"Image source: {path.name}\nContent type: {mime_type or 'image'}\nOCR: pending implementation."
    else:
        text = path.read_text(encoding="utf-8", errors="ignore")

    return {
        "text": text.strip(),
        "metadata": {
            "filename": path.name,
            "extension": suffix,
            "size_bytes": path.stat().st_size if path.exists() else 0,
        },
    }


def chunk_text(text: str, chunk_size: int = 1200) -> list[str]:
    cleaned = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    if not cleaned:
        return []

    chunks: list[str] = []
    buffer = ""
    for paragraph in cleaned.split("\n"):
        if len(buffer) + len(paragraph) + 1 > chunk_size and buffer:
            chunks.append(buffer.strip())
            buffer = ""
        buffer += paragraph + "\n"
    if buffer.strip():
        chunks.append(buffer.strip())
    return chunks


def _parse_csv(path: Path) -> str:
    lines: list[str] = []
    with path.open("r", encoding="utf-8-sig", errors="ignore", newline="") as handle:
        reader = csv.DictReader(handle)
        if reader.fieldnames:
            for row in reader:
                values = [f"{key}: {value}" for key, value in row.items() if value]
                lines.append(", ".join(values))
        else:
            handle.seek(0)
            raw_reader = csv.reader(handle)
            lines.extend(", ".join(cell for cell in row if cell) for row in raw_reader)
    return "\n".join(lines)


def _parse_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    lines: list[str] = []
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
        headers = [str(value).strip() if value is not None else "" for value in rows[0]]
        for row in rows[1:]:
            cells = ["" if value is None else str(value).strip() for value in row]
            if any(cells):
                if any(headers):
                    lines.append(", ".join(f"{headers[index]}: {cell}" for index, cell in enumerate(cells) if cell and index < len(headers)))
                else:
                    lines.append(", ".join(cell for cell in cells if cell))
    return "\n".join(lines)


def _parse_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _parse_docx(path: Path) -> str:
    from docx import Document

    document = Document(str(path))
    return "\n".join(paragraph.text for paragraph in document.paragraphs)
